"""Build assets/estimate_template.xlsx from the company reference workbook.

The template keeps the reference file's header rows 1-3 (fonts, theme fills,
borders, merges, column widths) and its Sheet2 rate table, and drops the
sample part rows, their images and the unrelated Sheet1.

Run once whenever the reference layout changes:
    python tools/make_template.py "ProCaliber Estimate Guide.xlsx"
"""

import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, os.pardir, "assets", "estimate_template.xlsx")


def build(reference_path, out_path=OUT):
    wb = openpyxl.load_workbook(reference_path)
    ws = wb["tube"]

    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 4:
            ws.unmerge_cells(str(rng))
    ws.delete_rows(4, ws.max_row)
    for row in [r for r in ws.row_dimensions if r >= 4]:
        del ws.row_dimensions[row]
    ws._images = []
    ws.conditional_formatting = type(ws.conditional_formatting)()
    ws.freeze_panes = "A4"
    # Cycle times are divided by 60 in the formulas, i.e. they are minutes.
    ws["AA2"] = "VMC Mins/piece"
    ws["AC2"] = "CNC Mins/piece"

    if "Sheet1" in wb.sheetnames:
        del wb["Sheet1"]

    rates = wb["Sheet2"]
    thin = Side(style="thin")
    for row in (1, 2):
        for col in ("A", "B"):
            cell = rates[f"{col}{row}"]
            cell.font = Font(name="Arial", size=11, bold=(col == "A"))
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center")
        rates[f"B{row}"].number_format = '"₹"\\ #,##0.00'
    rates.column_dimensions["A"].width = 30
    rates.column_dimensions["B"].width = 14

    ws.title = "ProCaliber Estimate"
    rates.title = "Machining cost"
    wb.active = 0
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    src = sys.argv[1] if len(sys.argv) > 1 else "ProCaliber Estimate Guide.xlsx"
    print("Wrote", os.path.normpath(build(src)))
