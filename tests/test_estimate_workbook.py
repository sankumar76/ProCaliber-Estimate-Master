import datetime
import os
import shutil
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), os.pardir))

import openpyxl
from PIL import Image

from core.estimate_workbook import (
    Dimensions, EstimateError, EstimateWorkbook, FileLockedError, PartInput,
    compute_costs, is_file_locked)

ROOT = os.path.join(os.path.dirname(__file__), os.pardir)
REFERENCE = os.path.join(ROOT, "ProCaliber Estimate Guide.xlsx")


def bush_part(**overrides):
    """Part 1 of the reference file (MS bush, 53 dia x 16 RM -> 48/38 x 12 FS)."""
    values = dict(
        material="MS (Mild Steel)", specific_gravity=7.85, batch_qty=12,
        rm_price_per_kg=69, date=datetime.date(2026, 9, 15),
        drawing_number="YP00013614-BSP-001", rev_no="-",
        description="Expansion Tank Inch BSP Bush", process_flow="CUT + TURN",
        rm=Dimensions(outer_dia=53, round_length=16),
        fs=Dimensions(outer_dia=48, inner_dia=38, round_length=12),
        cnc_min_per_piece=8, cnc_setting_hours=1.5)
    values.update(overrides)
    return PartInput(**values)


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_first_part_cells(self):
        book = EstimateWorkbook.new()
        self.assertEqual(book.add_part(bush_part()), 4)
        path = os.path.join(self.tmp, "a.xlsx")
        book.save(path)
        ws = openpyxl.load_workbook(path)["ProCaliber Estimate"]
        self.assertEqual(ws["A4"].value, 1)
        self.assertEqual(ws["I4"].value, "MS (Mild Steel)")
        self.assertEqual(ws["J4"].value, "MS (Mild Steel)")
        self.assertAlmostEqual(ws["U4"].value, 7.85e-6)
        self.assertEqual(ws["N4"].value, "RM")
        self.assertEqual(ws["N5"].value, "FS")
        self.assertIsNone(ws["O4"].value)  # unused dimensions stay blank
        self.assertEqual(ws["R5"].value, 48)
        self.assertEqual(ws["V4"].value,
                         "=(O4*P4*Q4*U4)+(3.14*(R4/2)*(R4/2)*T4*U4)-(3.14*(S4/2)*(S4/2)*T4*U4)")
        self.assertEqual(ws["W4"].value,
                         "=(O5*P5*Q5*U4)+(3.14*(R5/2)*(R5/2)*T5*U4)-(3.14*(S5/2)*(S5/2)*T5*U4)")
        self.assertEqual(ws["Y4"].value, "=V4*X4")
        self.assertEqual(ws["Z4"].value, 10)
        self.assertEqual(ws["AF4"].value, "=AD4*'Machining cost'!$B$1")
        self.assertEqual(ws["AG4"].value, "=AC4*('Machining cost'!$B$1/60)")
        self.assertEqual(ws["AI4"].value, "=AB4*'Machining cost'!$B$2")
        self.assertEqual(ws["AJ4"].value, "=AA4*('Machining cost'!$B$2/60)")
        self.assertEqual(ws["AK4"].value, "=SUM(AJ4,AG4,Z4,Y4)")
        self.assertEqual(ws["AL4"].value, "=AK4*0%")
        self.assertEqual(ws["AO4"].value, "=((V4-W4)*80%)*(X4/4)")
        self.assertEqual(ws["AP4"].value, "=(SUM(Y4,Z4,AE4:AN4)-AK4-AO4)*G4*115%")
        self.assertEqual(ws["AQ4"].value, "=(SUM(AK4,AL4:AN4)-AO4)*G4*115%")
        self.assertEqual(ws["AR4"].value, "=(AK4-AO4)*1.15")
        self.assertEqual(ws["AS4"].value, "=AP4+((H4-1)*AR4)")
        self.assertEqual(ws["AT4"].value, "=AS4/H4")
        self.assertIn("V4:V5", {str(r) for r in ws.merged_cells.ranges})
        self.assertNotIn("R4:R5", {str(r) for r in ws.merged_cells.ranges})

    def test_append_and_reopen(self):
        book = EstimateWorkbook.new()
        book.add_part(bush_part())
        self.assertEqual(book.add_part(bush_part(drawing_number="2")), 6)
        self.assertEqual(book.ws["AK6"].value, "=SUM(AJ6,AG6,Z6,Y6)")
        self.assertEqual(book.ws["AF6"].value, "=AD6*'Machining cost'!$B$1")
        path = os.path.join(self.tmp, "b.xlsx")
        book.save(path)
        reopened = EstimateWorkbook.open(path)
        self.assertEqual(reopened.add_part(bush_part(drawing_number="3")), 8)
        self.assertEqual(reopened.ws["A8"].value, 3)

    def test_open_reference_file_appends_after_last_part(self):
        book = EstimateWorkbook.open(REFERENCE)
        self.assertEqual(book.next_part_row(), 14)
        self.assertEqual(book.next_serial(), 6)
        self.assertEqual(book.rates, (300, 450))

    def test_old_sheet_names_are_migrated(self):
        book = EstimateWorkbook.open(REFERENCE)  # uses "tube" / "Sheet2"
        self.assertIn("ProCaliber Estimate", book.wb.sheetnames)
        self.assertIn("Machining cost", book.wb.sheetnames)
        self.assertNotIn("tube", book.wb.sheetnames)
        self.assertEqual(book.ws["AF6"].value, "=AD6*'Machining cost'!$B$1")
        self.assertEqual(book.ws["AJ12"].value, "=AA12*('Machining cost'!$B$2/60)")

    def test_save_refuses_file_open_elsewhere_and_leaves_it_intact(self):
        import ctypes
        path = os.path.join(self.tmp, "locked.xlsx")
        shutil.copy(REFERENCE, path)
        with open(path, "rb") as fh:
            before = fh.read()
        book = EstimateWorkbook.open(path)
        book.add_part(bush_part())
        kernel32 = ctypes.windll.kernel32
        kernel32.CreateFileW.restype = ctypes.c_void_p
        # Share-read-only handle: others may read but not write, like Excel.
        handle = kernel32.CreateFileW(path, 0x80000000, 1, None, 3, 0x80, None)
        try:
            self.assertTrue(is_file_locked(path))
            with self.assertRaises(FileLockedError):
                book.save()
        finally:
            kernel32.CloseHandle(ctypes.c_void_p(handle))
        with open(path, "rb") as fh:
            self.assertEqual(fh.read(), before)
        self.assertEqual([f for f in os.listdir(self.tmp) if f.startswith("~procaliber")], [])
        book.save()  # works once the other program lets go
        self.assertEqual(EstimateWorkbook.open(path).next_part_row(), 16)

    def test_save_keeps_backup_of_previous_version(self):
        old = os.environ.get("LOCALAPPDATA")
        os.environ["LOCALAPPDATA"] = self.tmp
        self.addCleanup(lambda: os.environ.__setitem__("LOCALAPPDATA", old) if old else None)
        path = os.path.join(self.tmp, "est.xlsx")
        book = EstimateWorkbook.new()
        book.add_part(bush_part())
        book.save(path)                      # new file: nothing to back up
        book.add_part(bush_part())
        book.save()                          # overwrite: previous version backed up
        backups = os.listdir(os.path.join(self.tmp, "ProCaliber Estimator", "Backups"))
        self.assertEqual(len(backups), 1)
        self.assertTrue(backups[0].startswith("est ("))

    def test_remove_last_part(self):
        img_path = os.path.join(self.tmp, "view.png")
        Image.new("RGB", (200, 150), "steelblue").save(img_path)
        book = EstimateWorkbook.new()
        book.add_part(bush_part(drawing_number="KEEP"))
        book.add_part(bush_part(drawing_number="UNDO", part_view_image=img_path))
        removed = book.remove_last_part()
        self.assertEqual(removed["drawing"], "UNDO")
        self.assertEqual(book.next_part_row(), 6)
        self.assertIsNone(book.ws["AT6"].value)
        self.assertEqual(book.ws._images, [])
        self.assertFalse(any(r.min_row >= 6 for r in book.ws.merged_cells.ranges))
        self.assertEqual(book.add_part(bush_part(drawing_number="AGAIN")), 6)
        path = os.path.join(self.tmp, "undo.xlsx")
        book.save(path)
        self.assertEqual([p["drawing"] for p in EstimateWorkbook.open(path).parts()],
                         ["KEEP", "AGAIN"])

    def test_rejects_wrong_layout(self):
        path = os.path.join(self.tmp, "bad.xlsx")
        wb = openpyxl.Workbook()
        wb.active.title = "ProCaliber Estimate"
        wb.save(path)
        with self.assertRaises(EstimateError):
            EstimateWorkbook.open(path)

    def test_rates_written_once_and_editable(self):
        book = EstimateWorkbook.new()
        self.assertEqual(book.rates, (300, 450))
        book.set_rates(350, 500)
        self.assertEqual(book.wb["Machining cost"]["B1"].value, 350)
        self.assertEqual(book.wb["Machining cost"]["B2"].value, 500)

    def test_percentages_and_image(self):
        img_path = os.path.join(self.tmp, "view.png")
        Image.new("RGB", (400, 300), "steelblue").save(img_path)
        book = EstimateWorkbook.new()
        book.add_part(bush_part(inspection_pct=5, part_view_image=img_path))
        self.assertEqual(book.ws["AL4"].value, "=AK4*5%")
        path = os.path.join(self.tmp, "c.xlsx")
        book.save(path)
        self.assertEqual(len(openpyxl.load_workbook(path)["ProCaliber Estimate"]._images), 1)

    def test_cycle_time_headers_say_minutes(self):
        self.assertEqual(EstimateWorkbook.new().ws["AA2"].value, "VMC Mins/piece")
        book = EstimateWorkbook.open(REFERENCE)  # reference still says Hours/piece
        self.assertEqual(book.ws["AA2"].value, "VMC Mins/piece")
        self.assertEqual(book.ws["AC2"].value, "CNC Mins/piece")

    def test_fs_size_optional(self):
        book = EstimateWorkbook.new()
        book.add_part(bush_part(fs=Dimensions()))
        for col in "OPQRST":
            self.assertIsNone(book.ws[f"{col}5"].value)
        self.assertEqual(compute_costs(bush_part(fs=Dimensions()), 300, 450)["finished_weight"], 0)

    def test_preview_matches_reference_numbers(self):
        c = compute_costs(bush_part(), 300, 450)
        rm = 3.14 * 26.5 ** 2 * 16 * 7.85e-6
        self.assertAlmostEqual(c["rm_weight"], rm)
        self.assertAlmostEqual(c["subtotal"], rm * 69 + 10 + 8 * 300 / 60)
        self.assertAlmostEqual(c["cost_per_piece"], c["total_cost"] / 12)


if __name__ == "__main__":
    unittest.main()
