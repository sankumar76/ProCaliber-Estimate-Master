"""Cost Estimator workbook logic (spec sections 5-8).

Produces / appends to an .xlsx in the "ProCaliber Estimate Guide" layout:
sheet ``ProCaliber Estimate`` (``tube`` in the reference file) holds one
row-pair per part (RM row n, FS row n+1, n even, starting at 4) with live
Excel formulas; sheet ``Machining cost`` (``Sheet2`` in the reference file)
holds the workbook-wide CNC (B1) and VMC (B2) machining rates.
"""

import datetime
import io
import os
import shutil
import sys
import tempfile
import zipfile
from dataclasses import dataclass, field

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.units import pixels_to_EMU

from core.materials import density_factor

SHEET = "ProCaliber Estimate"
RATES_SHEET = "Machining cost"
# Names used by the original reference file and by estimates saved before
# the rename; such files are migrated when opened.
LEGACY_SHEET = "tube"
LEGACY_RATES_SHEET = "Sheet2"
CNC_RATE_CELL = "B1"
VMC_RATE_CELL = "B2"
CNC_RATE_REF = f"'{RATES_SHEET}'!$B$1"
VMC_RATE_REF = f"'{RATES_SHEET}'!$B$2"
DEFAULT_CNC_RATE = 300.0
DEFAULT_VMC_RATE = 450.0
DEFAULT_CUTTING_COST = 10.0
FIRST_ROW = 4
ROW_HEIGHT = 54.75

CURRENCY = '"₹"\\ #,##0.00'

# Headers checked before appending to an existing file: cell -> expected text.
EXPECTED_HEADERS = {
    "A2": "S.no", "C2": "Drawing number", "H2": "Batch QTY",
    "N2": "Raw material Cost", "O3": "Length (mm)", "U3": "Density (Kg/m3)",
    "V3": "RM Weight/piece (kg)", "X3": "RM Price / KG",
    "AK2": "Subtotal: Cost / Piece", "AT2": "Cost / Piece",
}

# Cycle-time headers: the formulas divide these columns by 60, so they are
# minutes. Older files said "Hours/piece"; they are corrected on open.
CYCLE_TIME_HEADERS = {"AA2": "VMC Mins/piece", "AC2": "CNC Mins/piece"}

# Columns that are NOT merged across the RM/FS row pair (material size block).
UNMERGED_COLUMNS = {"N", "O", "P", "Q", "R", "S", "T"}
LAST_COLUMN = "AT"

NUMBER_FORMATS = {
    "A": "0", "B": "mm-dd-yy", "C": "0", "D": "0.000", "E": "0.000",
    "G": "0", "H": "0", "I": "0.000", "J": "0.000", "K": "0.000",
    "L": "0.000", "M": "0.000", "V": "0.000", "W": "0.000",
    "AA": "0.00", "AB": "0.00", "AC": "0.00", "AD": "0.00",
}
for _col in ("X", "Y", "Z", "AE", "AF", "AG", "AH", "AI", "AJ", "AK", "AL",
             "AM", "AN", "AO", "AP", "AQ", "AR", "AS", "AT"):
    NUMBER_FORMATS[_col] = CURRENCY

DIM_COLUMNS = ("O", "P", "Q", "R", "S", "T")

# Part-view picture box (column F is ~198 px wide, a row pair ~146 px tall).
IMAGE_MAX_W, IMAGE_MAX_H = 186, 136


class EstimateError(Exception):
    """A user-facing problem with an estimate workbook."""


class FileLockedError(EstimateError):
    """The target file is open in another program (usually Excel)."""

    def __init__(self, path):
        self.path = path
        super().__init__(
            f"“{os.path.basename(path)}” is open in another program "
            "(probably Excel), so it can't be saved right now.\n\n"
            "Close it there and click Save again. Your file has not been changed.")


def app_data_dir(*parts):
    """Per-user folder for backups and logs (created on demand)."""
    base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    folder = os.path.join(base, "ProCaliber Estimator", *parts)
    os.makedirs(folder, exist_ok=True)
    return folder


def is_file_locked(path):
    """True if another program (e.g. Excel) has the file open for editing."""
    if not os.path.exists(path):
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True
    except OSError:
        return False


BACKUPS_TO_KEEP = 50


def backup_file(path):
    """Copy ``path`` into the backups folder with a timestamp; best effort."""
    try:
        folder = app_data_dir("Backups")
        stem, ext = os.path.splitext(os.path.basename(path))
        stamp = datetime.datetime.now().strftime("%Y-%m-%d %H%M%S")
        shutil.copy2(path, os.path.join(folder, f"{stem} ({stamp}){ext}"))
        old = sorted((os.path.join(folder, f) for f in os.listdir(folder)),
                     key=os.path.getmtime)
        for stale in old[:-BACKUPS_TO_KEEP]:
            os.remove(stale)
    except OSError:
        pass


def resource_path(*parts):
    """Locate bundled files both from source and inside a PyInstaller exe."""
    base = getattr(sys, "_MEIPASS", os.path.join(os.path.dirname(__file__), os.pardir))
    return os.path.normpath(os.path.join(base, *parts))


TEMPLATE_PATH = resource_path("assets", "estimate_template.xlsx")


@dataclass
class Dimensions:
    """Material size block: rectangular (O,P,Q) and/or round (R,S,T), mm."""
    length: float = None        # O
    width: float = None         # P
    thickness: float = None     # Q
    outer_dia: float = None     # R
    inner_dia: float = None     # S
    round_length: float = None  # T

    def values(self):
        return (self.length, self.width, self.thickness,
                self.outer_dia, self.inner_dia, self.round_length)

    def is_empty(self):
        return all(v is None for v in self.values())

    def describe(self):
        """Short human-readable size, e.g. 'Ø53 × 16 mm' or '20 × 20 × 90 mm'."""
        parts = []
        if self.length is not None:
            parts.append(" × ".join(f"{v:g}" for v in (self.length, self.width, self.thickness)
                                        if v is not None) + " mm")
        if self.outer_dia is not None:
            text = f"Ø{self.outer_dia:g}"
            if self.inner_dia:
                text += f" / Ø{self.inner_dia:g} ID"
            if self.round_length is not None:
                text += f" × {self.round_length:g}"
            parts.append(text + " mm")
        return ", ".join(parts) or "no size"

    def weight(self, density):
        """Same arithmetic as the V/W formulas (blank = 0, pi = 3.14)."""
        o, p, q, r, s, t = (v or 0 for v in self.values())
        return (o * p * q * density
                + 3.14 * (r / 2) * (r / 2) * t * density
                - 3.14 * (s / 2) * (s / 2) * t * density)


@dataclass
class PartInput:
    material: str
    specific_gravity: float
    batch_qty: int
    rm_price_per_kg: float
    rm: Dimensions = field(default_factory=Dimensions)
    fs: Dimensions = field(default_factory=Dimensions)
    date: datetime.date = field(default_factory=datetime.date.today)
    drawing_number: str = ""
    rev_no: str = ""
    description: str = ""
    part_view_image: str = None   # path to a picture file, optional
    qty_per_frame: float = 1
    process_flow: str = ""
    coating: str = ""
    remarks: str = ""
    cutting_cost: float = DEFAULT_CUTTING_COST
    vmc_min_per_piece: float = 0      # AA (formula divides by 60)
    vmc_setting_hours: float = 0      # AB
    cnc_min_per_piece: float = 0      # AC (formula divides by 60)
    cnc_setting_hours: float = 0      # AD
    inspection_pct: float = 0         # AL
    transport_pct: float = 0          # AM
    packing_pct: float = 0            # AN


def compute_costs(part, cnc_rate, vmc_rate):
    """Python mirror of the row formulas, used for the on-screen preview."""
    u = density_factor(part.specific_gravity)
    v = part.rm.weight(u)
    w = part.fs.weight(u)
    y = v * part.rm_price_per_kg
    z = part.cutting_cost
    af = part.cnc_setting_hours * cnc_rate
    ag = part.cnc_min_per_piece * cnc_rate / 60
    ai = part.vmc_setting_hours * vmc_rate
    aj = part.vmc_min_per_piece * vmc_rate / 60
    ak = aj + ag + z + y
    al = ak * part.inspection_pct / 100
    am = ak * part.transport_pct / 100
    an = ak * part.packing_pct / 100
    ao = (v - w) * 0.8 * (part.rm_price_per_kg / 4)
    g = part.qty_per_frame
    ap = (y + z + af + ag + ai + aj + ak + al + am + an - ak - ao) * g * 1.15
    aq = (ak + al + am + an - ao) * g * 1.15
    ar = (ak - ao) * 1.15
    h = part.batch_qty
    as_ = ap + (h - 1) * ar
    at = as_ / h if h else None
    return {
        "density": u, "rm_weight": v, "finished_weight": w, "rm_cost": y,
        "cutting_cost": z, "cnc_setting_cost": af, "cnc_machining_cost": ag,
        "vmc_setting_cost": ai, "vmc_machining_cost": aj, "subtotal": ak,
        "inspection": al, "transport": am, "packing": an, "scrap_return": ao,
        "develop_cost": ap, "after_first_cost": aq, "prod_cost": ar,
        "total_cost": as_, "cost_per_piece": at,
    }


def row_formulas(n, inspection_pct=0, transport_pct=0, packing_pct=0):
    """Formula strings for the part whose RM row is ``n``."""
    m = n + 1

    def weight(r):
        return (f"=(O{r}*P{r}*Q{r}*U{n})+(3.14*(R{r}/2)*(R{r}/2)*T{r}*U{n})"
                f"-(3.14*(S{r}/2)*(S{r}/2)*T{r}*U{n})")

    return {
        "V": weight(n),
        "W": weight(m),
        "Y": f"=V{n}*X{n}",
        "AF": f"=AD{n}*{CNC_RATE_REF}",
        "AG": f"=AC{n}*({CNC_RATE_REF}/60)",
        "AI": f"=AB{n}*{VMC_RATE_REF}",
        "AJ": f"=AA{n}*({VMC_RATE_REF}/60)",
        "AK": f"=SUM(AJ{n},AG{n},Z{n},Y{n})",
        "AL": f"=AK{n}*{_pct(inspection_pct)}%",
        "AM": f"=AK{n}*{_pct(transport_pct)}%",
        "AN": f"=AK{n}*{_pct(packing_pct)}%",
        "AO": f"=((V{n}-W{n})*80%)*(X{n}/4)",
        "AP": f"=(SUM(Y{n},Z{n},AE{n}:AN{n})-AK{n}-AO{n})*G{n}*115%",
        "AQ": f"=(SUM(AK{n},AL{n}:AN{n})-AO{n})*G{n}*115%",
        "AR": f"=(AK{n}-AO{n})*1.15",
        "AS": f"=AP{n}+((H{n}-1)*AR{n})",
        "AT": f"=AS{n}/H{n}",
    }


def _pct(value):
    return f"{value:g}"


def _normalise(text):
    return " ".join(str(text or "").split()).lower()


class EstimateWorkbook:
    def __init__(self, wb, path=None):
        self.wb = wb
        self.path = path
        self.dirty = False

    # ---- create / open / save ---------------------------------------------

    @classmethod
    def new(cls):
        if not os.path.exists(TEMPLATE_PATH):
            raise EstimateError(f"Estimate template is missing:\n{TEMPLATE_PATH}")
        return cls(openpyxl.load_workbook(TEMPLATE_PATH))

    @classmethod
    def open(cls, path):
        try:
            wb = openpyxl.load_workbook(path)
        except PermissionError:
            raise EstimateError("The file is locked. Close it in Excel and try again.")
        except Exception as exc:
            raise EstimateError(f"Could not read this file as an Excel workbook.\n\n{exc}")
        legacy = SHEET not in wb.sheetnames and LEGACY_SHEET in wb.sheetnames
        name = LEGACY_SHEET if legacy else SHEET
        if name not in wb.sheetnames:
            raise EstimateError(
                f'This workbook has no "{SHEET}" sheet, so it is not a '
                "ProCaliber estimate.")
        ws = wb[name]
        wrong = [f"{cell} (expected \"{text}\")"
                 for cell, text in EXPECTED_HEADERS.items()
                 if _normalise(ws[cell].value) != _normalise(text)]
        if wrong:
            raise EstimateError(
                "The header layout does not match the ProCaliber estimate "
                "format. Unexpected cells:\n  " + "\n  ".join(wrong))
        book = cls(wb, path)
        if legacy:
            book._migrate_sheet_names()
        book._ensure_rates_sheet()
        book._fix_cycle_time_headers()
        # Format upgrades alone aren't user changes; they are written with the
        # next save, so opening an older file shouldn't prompt to save it.
        book.dirty = False
        return book

    def save(self, path=None):
        """Save safely: write a temporary file next to the target, check it,
        back up the previous version, then swap it into place. The existing
        file is never touched until a complete new copy is ready."""
        path = path or self.path
        if not path:
            raise EstimateError("Choose where to save the estimate first.")
        path = os.path.abspath(path)
        if is_file_locked(path):
            raise FileLockedError(path)
        folder = os.path.dirname(path)
        try:
            fd, tmp = tempfile.mkstemp(prefix="~procaliber-", suffix=".xlsx", dir=folder)
            os.close(fd)
        except OSError as exc:
            raise EstimateError(f"Can't write to the folder\n{folder}\n\n{exc}")
        try:
            self.wb.save(tmp)
            with zipfile.ZipFile(tmp) as archive:
                if archive.testzip() is not None:
                    raise EstimateError("The saved copy failed its integrity check. "
                                        "Nothing was changed; please try again.")
            if os.path.exists(path):
                backup_file(path)
            os.replace(tmp, path)
        except PermissionError:
            raise FileLockedError(path)
        except EstimateError:
            raise
        except Exception as exc:
            raise EstimateError(f"Could not save to\n{path}\n\n{exc}\n\n"
                                "The existing file was not changed.")
        finally:
            if os.path.exists(tmp):
                try:
                    os.remove(tmp)
                except OSError:
                    pass
        self.path = path
        self.dirty = False

    def _migrate_sheet_names(self):
        """Rename tube/Sheet2 to the current names and repoint the formulas."""
        self.wb[LEGACY_SHEET].title = SHEET
        if LEGACY_RATES_SHEET in self.wb.sheetnames and RATES_SHEET not in self.wb.sheetnames:
            self.wb[LEGACY_RATES_SHEET].title = RATES_SHEET
            old_ref = f"{LEGACY_RATES_SHEET}!"
            new_ref = f"'{RATES_SHEET}'!"
            for row in self.ws.iter_rows(min_row=FIRST_ROW):
                for cell in row:
                    if isinstance(cell.value, str) and old_ref in cell.value:
                        cell.value = cell.value.replace(old_ref, new_ref)
        self.dirty = True

    def _fix_cycle_time_headers(self):
        for cell, text in CYCLE_TIME_HEADERS.items():
            if self.ws[cell].value != text:
                self.ws[cell] = text
                self.dirty = True

    # ---- workbook-wide machining rates -------------------------------------

    def _ensure_rates_sheet(self):
        if RATES_SHEET not in self.wb.sheetnames:
            rates = self.wb.create_sheet(RATES_SHEET)
            rates["A1"] = "CNC Machining Cost / Hour"
            rates["A2"] = "VMC Machining Cost / Hour"
            rates[CNC_RATE_CELL] = DEFAULT_CNC_RATE
            rates[VMC_RATE_CELL] = DEFAULT_VMC_RATE
            self.dirty = True

    @property
    def rates(self):
        """(cnc_rate, vmc_rate) currently stored in the workbook."""
        rates = self.wb[RATES_SHEET]
        return (_as_number(rates[CNC_RATE_CELL].value, DEFAULT_CNC_RATE),
                _as_number(rates[VMC_RATE_CELL].value, DEFAULT_VMC_RATE))

    def set_rates(self, cnc_rate, vmc_rate):
        if (cnc_rate, vmc_rate) == self.rates:
            return
        rates = self.wb[RATES_SHEET]
        rates[CNC_RATE_CELL] = cnc_rate
        rates[VMC_RATE_CELL] = vmc_rate
        self.dirty = True

    # ---- parts ------------------------------------------------------------

    @property
    def ws(self):
        return self.wb[SHEET]

    def _part_rows(self):
        """RM rows that hold a numeric S.no in column A."""
        rows = []
        for row in range(FIRST_ROW, self.ws.max_row + 1):
            if isinstance(self.ws.cell(row, 1).value, (int, float)):
                rows.append(row)
        return rows

    def next_part_row(self):
        rows = self._part_rows()
        if not rows:
            return FIRST_ROW
        last = rows[-1]
        return last + 2 if last % 2 == 0 else last + 1

    def next_serial(self):
        rows = self._part_rows()
        return int(max(self.ws.cell(r, 1).value for r in rows)) + 1 if rows else 1

    def parts(self):
        """Summary of each part for the on-screen list."""
        out = []
        for row in self._part_rows():
            cell = lambda col: self.ws[f"{col}{row}"].value
            out.append({
                "row": row, "sno": cell("A"), "drawing": cell("C"),
                "description": cell("E"), "material": cell("J"),
                "batch_qty": cell("H"),
            })
        return out

    def add_part(self, part):
        """Write ``part`` into the next free row pair; returns its RM row."""
        n = self.next_part_row()
        m = n + 1
        ws = self.ws
        values = {
            "A": self.next_serial(),
            "B": datetime.datetime.combine(part.date, datetime.time()),
            "C": part.drawing_number, "D": part.rev_no, "E": part.description,
            "G": part.qty_per_frame, "H": part.batch_qty,
            "I": part.material, "J": part.material,
            "K": part.process_flow, "L": part.coating, "M": part.remarks,
            "N": "RM",
            "U": density_factor(part.specific_gravity),
            "X": part.rm_price_per_kg, "Z": part.cutting_cost,
            "AA": part.vmc_min_per_piece, "AB": part.vmc_setting_hours,
            "AC": part.cnc_min_per_piece, "AD": part.cnc_setting_hours,
        }
        values.update(row_formulas(n, part.inspection_pct,
                                   part.transport_pct, part.packing_pct))

        self._style_row_pair(n)
        for col, value in values.items():
            if value == "" or value is None:
                continue
            ws[f"{col}{n}"] = value
        ws[f"N{m}"] = "FS"
        for col, rm_value, fs_value in zip(DIM_COLUMNS, part.rm.values(), part.fs.values()):
            if rm_value is not None:
                ws[f"{col}{n}"] = rm_value
            if fs_value is not None:
                ws[f"{col}{m}"] = fs_value

        if part.part_view_image:
            self._add_image(part.part_view_image, n)
        self.dirty = True
        return n

    def remove_last_part(self):
        """Remove the last part's row pair (and its picture); returns its summary."""
        rows = self._part_rows()
        if not rows:
            raise EstimateError("There are no parts in this spreadsheet to remove.")
        n = rows[-1]
        removed = self.parts()[-1]
        ws = self.ws
        for rng in list(ws.merged_cells.ranges):
            if rng.min_row >= n and rng.max_row <= n + 1:
                ws.unmerge_cells(str(rng))
        last = column_index_from_string(LAST_COLUMN)
        for row in (n, n + 1):
            for idx in range(1, last + 1):
                cell = ws.cell(row, idx)
                cell.value = None
                cell.style = "Normal"
            ws.row_dimensions[row].height = None
        ws._images = [img for img in ws._images
                      if not n - 1 <= img.anchor._from.row <= n]
        self.dirty = True
        return removed

    def _style_row_pair(self, n):
        ws = self.ws
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        font = Font(name="Arial", size=11)
        align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        last = column_index_from_string(LAST_COLUMN)
        for row in (n, n + 1):
            ws.row_dimensions[row].height = ROW_HEIGHT
            for idx in range(1, last + 1):
                cell = ws.cell(row, idx)
                cell.font = font
                cell.border = border
                cell.alignment = align
                fmt = NUMBER_FORMATS.get(get_column_letter(idx))
                if fmt:
                    cell.number_format = fmt
        for idx in range(1, last + 1):
            col = get_column_letter(idx)
            if col not in UNMERGED_COLUMNS:
                ws.merge_cells(f"{col}{n}:{col}{n + 1}")

    def _add_image(self, path, n):
        from PIL import Image as PILImage

        try:
            with PILImage.open(path) as src:
                src.load()
                pic = src.convert("RGBA") if src.mode not in ("RGB", "RGBA") else src.copy()
        except Exception as exc:
            raise EstimateError(f"Could not read the part view picture:\n{path}\n\n{exc}")
        scale = min(IMAGE_MAX_W / pic.width, IMAGE_MAX_H / pic.height, 1.0)
        disp_w, disp_h = max(int(pic.width * scale), 1), max(int(pic.height * scale), 1)
        # Keep 2x resolution so the picture stays sharp when zoomed in Excel.
        pic.thumbnail((disp_w * 2, disp_h * 2))
        buf = io.BytesIO()
        pic.save(buf, format="PNG")
        buf.seek(0)
        img = XLImage(buf)
        img.width, img.height = disp_w, disp_h
        col_px = int((self.ws.column_dimensions["F"].width or 27.55) * 7 + 5)
        pair_px = int(ROW_HEIGHT * 2 * 96 / 72)
        marker = AnchorMarker(
            col=column_index_from_string("F") - 1, row=n - 1,
            colOff=pixels_to_EMU(max((col_px - disp_w) // 2, 0)),
            rowOff=pixels_to_EMU(max((pair_px - disp_h) // 2, 0)))
        img.anchor = OneCellAnchor(
            _from=marker,
            ext=XDRPositiveSize2D(pixels_to_EMU(disp_w), pixels_to_EMU(disp_h)))
        self.ws.add_image(img)


def _as_number(value, default):
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).lstrip("="))
    except (TypeError, ValueError):
        return default
